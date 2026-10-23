"""Convierte la plantilla de datos (xlsx) en knowledge.json para el motor de inferencia.
Uso: python3 cargar_conocimiento.py plantilla_datos_40_paises.xlsx knowledge.json
Lee tambien reglas.json y preguntas.json (misma carpeta), calcula la afinidad 1-5 de cada pais
por criterio (quintiles o mapa fijo) y lo junta todo en knowledge.json.
Regla: un dato sin valor, sin URL de fuente o sin fecha se guarda como "SIN DATO" y se avisa.
"""
import sys, json, os
from openpyxl import load_workbook

NUMERICOS = {"coste_diario","horas_vuelo","patrimonio_cultural","naturaleza","gastronomia"}
SIN = "SIN DATO"
NIVEL_TRAMITE = {"Sin trámite previo": 0, "Autorización electrónica": 1, "Visado a la llegada": 2, "Visado electrónico": 2, "Visado consular": 3}

def num(v):
    try: return float(str(v).replace(",", "."))
    except (TypeError, ValueError): return None

def cargar_json(nombre):
    ruta = os.path.join(os.path.dirname(os.path.abspath(__file__)), nombre)
    with open(ruta, encoding="utf-8") as f: return json.load(f)

def calcular_afinidades(paises, criterios):
    """Quintil por rango percentil entre los paises con dato (empates a rango medio) o mapa fijo."""
    for cr in criterios:
        if cr["metodo"] == "dinamico": continue
        k = cr["atributo"]
        vals = {p["nombre"]: p["atributos"].get(k, {}).get("valor") for p in paises}
        if cr["metodo"] == "mapa":
            for p in paises: p.setdefault("afinidad", {})[cr["id"]] = cr["mapa"].get(vals[p["nombre"]])
        else:
            nums = [v for v in vals.values() if isinstance(v, (int, float))]
            n = len(nums)
            for p in paises:
                v = vals[p["nombre"]]
                if not isinstance(v, (int, float)) or n == 0: a = None
                else:
                    menos = sum(1 for x in nums if x < v); iguales = sum(1 for x in nums if x == v)
                    a = min(5, 1 + int(((menos + 0.5 * iguales) / n) * 5))
                    if cr.get("invertir"): a = 6 - a
                p.setdefault("afinidad", {})[cr["id"]] = a

def main(src, dst):
    wb = load_workbook(src, data_only=True)
    paises = {}
    for fila in list(wb["Países"].iter_rows(min_row=2, values_only=True)):
        n, p, z, s = fila[:4]
        if p: paises[p] = {"id": int(n), "nombre": p, "zona": z, "subzona": s, "atributos": {}, "clima": None}
    avisos = []
    for fila in wb["Datos"].iter_rows(min_row=2, values_only=True):
        p, k, _n, v, u, url, fecha, ver, obs = fila[:9]
        periodo, metodo, calidad = (list(fila[9:12]) + [None]*3)[:3]
        if p not in paises: continue
        vacio = v in (None, "") or not url or not fecha
        if vacio:
            avisos.append(f"{p} / {k}: falta valor, URL o fecha"); val = SIN
        else:
            val = num(v) if k in NUMERICOS else str(v)
            if k in NUMERICOS and val is None: avisos.append(f"{p} / {k}: valor no numérico ({v})"); val = SIN
        paises[p]["atributos"][k] = {"valor": val, "unidad": u, "fuente": url, "fecha": str(fecha) if fecha else None,
                                       "verificado": ver == "Sí", "observaciones": obs,
                                       "periodo_referencia": periodo, "metodo_obtencion": metodo, "calidad": calidad}
    for fila in wb["Clima"].iter_rows(min_row=2, values_only=True):
        p = fila[0]
        if p not in paises: continue
        t, pr = [num(x) for x in fila[1:13]], [num(x) for x in fila[13:25]]
        url, fecha, ver, obs = fila[25], fila[26], fila[27], fila[28]
        if None in t or None in pr or not url or not fecha:
            avisos.append(f"{p} / clima: faltan meses, URL o fecha"); t = pr = None
        paises[p]["clima"] = {"temperatura": t, "precipitacion": pr, "fuente": url,
                               "fecha": str(fecha) if fecha else None, "verificado": ver == "Sí", "observaciones": obs}
    for p in paises.values():
        e = p["atributos"].get("entrada_espanoles", {}).get("valor")
        p["nivel_tramite"] = NIVEL_TRAMITE.get(e)
        if p["nivel_tramite"] is None: avisos.append(f"{p['nombre']} / nivel_tramite: entrada sin valor válido")
    reglas = cargar_json("reglas.json"); preguntas = cargar_json("preguntas.json")
    lista = list(paises.values())
    calcular_afinidades(lista, reglas["criterios"])
    kb = {"version": 1, "criterios": reglas["criterios"], "paises": lista,
          "reglas": reglas["reglas"], "preguntas": preguntas["preguntas"]}
    with open(dst, "w", encoding="utf-8") as f: json.dump(kb, f, ensure_ascii=False, indent=2)
    print(f"{len(paises)} países -> {dst}. Avisos: {len(avisos)}")
    for a in avisos[:15]: print(" -", a)
    if len(avisos) > 15: print(f" ... y {len(avisos)-15} más")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
